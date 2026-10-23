#!/usr/bin/env python3
"""Extrait du classeur maître (recalculé) les chiffres affichés dans l'appli.

Usage : python3 ledger_summary.py CLASSEUR.xlsx SORTIE.json
Le classeur doit contenir les valeurs calculées (sortie de ledger_sync.py ou d'un recalcul LibreOffice).
SORTIE.json est le document ledger/summary de l'appli.
"""
import datetime as dt
import json
import sys

import openpyxl


def num(v):
    return round(float(v), 2) if isinstance(v, (int, float)) else 0.0


def main(src, dst):
    wb = openpyxl.load_workbook(src, data_only=True)
    J, B = wb['Journal'], wb['Balance']

    accounts = []
    for r in range(5, B.max_row + 1):
        code = B.cell(r, 1).value
        if not code or not str(code).isdigit():
            continue
        deb, cre = num(B.cell(r, 4).value), num(B.cell(r, 5).value)
        if deb or cre:
            accounts.append({'code': str(code), 'label': B.cell(r, 2).value, 'debit': deb, 'credit': cre,
                             'balance': round(deb - cre, 2)})
    bal = {a['code']: a['balance'] for a in accounts}

    lines, months = [], {}
    for r in range(5, 1005):
        no = J.cell(r, 1).value
        if no in (None, ''):
            continue
        date = J.cell(r, 3).value
        iso = date.date().isoformat() if isinstance(date, dt.datetime) else str(date or '')[:10]
        acc = str(J.cell(r, 7).value or '')
        deb, cre = num(J.cell(r, 10).value), num(J.cell(r, 11).value)
        lines.append({'no': int(no), 'journal': J.cell(r, 2).value, 'date': iso, 'piece': str(J.cell(r, 5).value or ''),
                      'ref': J.cell(r, 6).value or '', 'account': acc, 'accountLabel': J.cell(r, 8).value or '',
                      'label': J.cell(r, 9).value or '', 'debit': deb, 'credit': cre})
        if acc[:1] in ('6', '7') and iso[:7]:
            m = months.setdefault(iso[:7], {'charges': 0.0, 'produits': 0.0})
            if acc[0] == '6':
                m['charges'] = round(m['charges'] + deb - cre, 2)
            else:
                m['produits'] = round(m['produits'] + cre - deb, 2)

    # Dotations prévues par année (onglet Immobilisations, colonnes d'années en ligne 4)
    I = wb['Immobilisations']
    depreciation = {}
    total_row = next((r for r in range(5, I.max_row + 1) if str(I.cell(r, 1).value or '').strip().upper() == 'TOTAL'), None)
    for c in range(1, I.max_column + 1):
        y = I.cell(4, c).value
        if total_row and (isinstance(y, (int, float)) or (isinstance(y, str) and y.isdigit())):
            depreciation[str(int(y))] = num(I.cell(total_row, c).value)

    charges = round(sum(a['balance'] for a in accounts if a['code'][0] == '6'), 2)
    produits = round(-sum(a['balance'] for a in accounts if a['code'][0] == '7'), 2)
    exercise = wb['Paramètres']['C10'].value
    out = {
        'generatedAt': dt.datetime.now(dt.timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'),
        'exerciseEnd': exercise.date().isoformat() if isinstance(exercise, dt.datetime) else '',
        'produits': produits, 'charges': charges, 'result': round(produits - charges, 2),
        'bank': bal.get('512100', 0.0),
        'loanOutstanding': round(-bal.get('164100', 0.0), 2),
        'owedToRegine': round(-bal.get('467600', 0.0), 2),
        'ownerAccount': round(-bal.get('108000', 0.0), 2),
        'fixedAssets': round(sum(v for k, v in bal.items() if k.startswith('21')), 2),
        'journalGap': J['J1007'].value,
        'bankGap': J['J1011'].value,
        'depreciation': depreciation,
        'accounts': accounts,
        'months': months,
        'lines': lines,
    }
    json.dump(out, open(dst, 'w', encoding='utf-8'), ensure_ascii=False)
    print(json.dumps({k: out[k] for k in ('produits', 'charges', 'result', 'bank', 'loanOutstanding', 'owedToRegine',
                                          'ownerAccount', 'fixedAssets', 'journalGap', 'bankGap')}, ensure_ascii=False),
          len(lines), 'lignes')


if __name__ == '__main__':
    main(*sys.argv[1:3])
